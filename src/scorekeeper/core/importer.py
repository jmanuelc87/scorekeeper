"""Parse a conversation ``.xlsx`` file into raw message rows.

The source spreadsheet stores a conversation as **one message per row** (a user
or model message with, optionally, a turn number). This module reads such a
file with ``openpyxl`` and projects it into a JSON-serializable list of message
dicts with the keys ``turn``, ``role`` and ``content`` (plus an optional
``retrieved_context_source`` — the **raw** ``retrieved_context`` cell, a ranked list
of source references the retrieval pipeline later fetches/extracts — and/or
``expected_output`` when the sheet carries those columns) — the raw conversation
shape that ``ScenarioResult.raw_conversation`` is documented to hold. Turn rows
for evaluation are derived from this projection in a later step; this module
only parses, and does **not** interpret the ``retrieved_context`` cell.

Header names are matched case-insensitively against a small alias table so the
parser tolerates either English or Spanish files (conversation content is in
Spanish). Callers with an unusual layout can pass an explicit ``columns``
mapping to override detection.

The tail of that projection — role aliasing and turn numbering — is exposed on its
own as :func:`normalize_messages`, because the spreadsheet is not the only source
of a conversation: the browser extension scrapes messages off a chat UI and posts
them to ``POST /captures``, which normalizes them the same way so both ingestion
paths produce identical rows.
"""

from __future__ import annotations

import json
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

# Header aliases (normalized: stripped + lowercased) -> canonical field name.
_HEADER_ALIASES: dict[str, str] = {
    "turn": "turn",
    "turno": "turn",
    "#": "turn",
    "role": "role",
    "rol": "role",
    "autor": "role",
    "content": "content",
    "contenido": "content",
    "mensaje": "content",
    "texto": "content",
    "retrieved context": "retrieved_context",
    "retrieved_context": "retrieved_context",
    "retrieval context": "retrieved_context",
    "retrieval_context": "retrieved_context",
    "contexto recuperado": "retrieved_context",
    "contexto": "retrieved_context",
    "expected output": "expected_output",
    "expected_output": "expected_output",
    "expected": "expected_output",
    "reference": "expected_output",
    "respuesta esperada": "expected_output",
    "salida esperada": "expected_output",
    "referencia": "expected_output",
    "esperado": "expected_output",
}

# Role value aliases (normalized) -> canonical role.
_ROLE_ALIASES: dict[str, str] = {
    "user": "user",
    "usuario": "user",
    "model": "model",
    "modelo": "model",
    "assistant": "model",
    "asistente": "model",
}


def _normalize(value: Any) -> str:
    """Strip and lowercase a cell/header value; ``None`` becomes ``""``."""
    if value is None:
        return ""
    return str(value).strip().lower()


def _map_columns(header: list[Any], override: dict[str, str] | None) -> dict[str, int]:
    """Map canonical field names to their 0-based column index in ``header``.

    ``override`` maps a canonical field (``turn``/``role``/``content``) to the
    header label as it literally appears in the sheet; it wins over aliases.
    Raises ``ValueError`` if ``role`` or ``content`` cannot be located.
    """
    normalized = [_normalize(cell) for cell in header]
    columns: dict[str, int] = {}

    if override:
        for field, label in override.items():
            target = _normalize(label)
            for idx, name in enumerate(normalized):
                if name == target:
                    columns[field] = idx
                    break

    for idx, name in enumerate(normalized):
        field = _HEADER_ALIASES.get(name)
        if field is not None and field not in columns:
            columns[field] = idx

    missing = [field for field in ("role", "content") if field not in columns]
    if missing:
        raise ValueError(
            f"No se encontraron columnas requeridas {missing} en el encabezado "
            f"{[c for c in header]!r}."
        )
    return columns


def normalize_messages(
    rows: Iterable[dict[str, Any]],
    *,
    derive_turns: bool | None = None,
) -> list[dict[str, Any]]:
    """Normalize raw message dicts into the canonical conversation shape.

    The shared tail of every ingestion path: the ``.xlsx`` parser feeds it one dict
    per sheet row, and the browser extension's captures (``POST /captures``) feed it
    one dict per scraped chat bubble, so both reach the database through the same
    role aliasing and turn numbering.

    Each input row may carry ``turn``, ``role``, ``content``, ``retrieved_context``
    and ``expected_output``; the optional keys are carried through only when the
    input actually has them. Rows with neither a role nor content are dropped.

    ``derive_turns`` forces turn numbering: ``True`` always derives (each ``user``
    message following a non-user message starts a new turn, so a user+model pair
    shares a number), ``False`` always reads the row's ``turn`` (blank → ``1``), and
    the default ``None`` derives only when no row carries a turn value at all.
    """
    rows = list(rows)
    if derive_turns is None:
        derive_turns = not any(str(row.get("turn") or "").strip() for row in rows)

    messages: list[dict[str, Any]] = []
    derived_turn = 0
    prev_role: str | None = None

    for row in rows:
        content = _stringify(row.get("content")).strip()
        role_raw = _stringify(row.get("role")).strip()
        if not content and not role_raw:
            continue  # Fully-empty row.

        role = _ROLE_ALIASES.get(role_raw.lower(), role_raw.lower())

        if derive_turns:
            if role == "user" and prev_role != "user":
                derived_turn += 1
            elif derived_turn == 0:
                derived_turn = 1
            turn = derived_turn
        else:
            turn = _coerce_turn(_stringify(row.get("turn")))

        message: dict[str, Any] = {"turn": turn, "role": role, "content": content}
        for key in ("retrieved_context", "expected_output"):
            if key in row:
                message[key] = _stringify(row[key]).strip()
        messages.append(message)
        prev_role = role

    return messages


def parse_conversation(
    path: str | Path,
    *,
    columns: dict[str, str] | None = None,
) -> list[dict[str, Any]]:
    """Parse an ``.xlsx`` conversation file into raw message dicts.

    Returns a list of ``{"turn": int, "role": str, "content": str}`` in sheet
    order; each dict also carries ``retrieved_context_source`` (the raw context
    cell string, interpreted later by the retrieval pipeline) and/or
    ``expected_output`` (a string) when the sheet has those columns. Roles are
    normalized to canonical values
    (``user``/``model``) when recognized, otherwise passed through normalized.
    When the sheet has no turn column, turn numbers are derived: each ``user``
    message that follows a non-user message starts a new turn, so a user+model
    pair shares one number.

    ``columns`` optionally maps a canonical field (``turn``/``role``/``content``/
    ``retrieved_context``/``expected_output``) to the exact header label in the
    sheet, overriding alias detection.

    Raises ``FileNotFoundError`` if ``path`` does not exist and ``ValueError``
    if the required ``role``/``content`` columns cannot be found.
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"No existe el archivo de conversación: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            return []  # No worksheets in the workbook.
        rows = sheet.iter_rows(values_only=True)
        try:
            header = list(next(rows))
        except StopIteration:
            return []  # Empty sheet — no header, no messages.

        col = _map_columns(header, columns)
        has_turn_col = "turn" in col

        # Project the sheet onto the raw dicts normalize_messages consumes: only the
        # columns the sheet actually has, so absent optional columns stay absent.
        raw_rows: list[dict[str, Any]] = []
        for raw in rows:
            row = list(raw)
            entry = {
                "role": _cell(row, col["role"]),
                "content": _cell(row, col["content"]),
            }
            for field in ("turn", "retrieved_context", "expected_output"):
                if field in col:
                    entry[field] = _cell(row, col[field])
            raw_rows.append(entry)

        messages = normalize_messages(raw_rows, derive_turns=not has_turn_col)
        # The .xlsx path exposes the raw context cell as ``retrieved_context_source``
        # (the retrieval pipeline interprets it later); ``normalize_messages`` keeps the
        # generic ``retrieved_context`` key the browser-capture path relies on, so rename
        # it here to match the DB column.
        for message in messages:
            if "retrieved_context" in message:
                message["retrieved_context_source"] = message.pop("retrieved_context")
        return messages
    finally:
        workbook.close()


def parse_conversation_json(
    path: str | Path,
    *,
    columns: dict[str, str] | None = None,
) -> str:
    """Parse ``path`` and return the raw conversation as a JSON string.

    Uses ``ensure_ascii=False`` so Spanish accents are preserved verbatim.
    """
    return json.dumps(
        parse_conversation(path, columns=columns), ensure_ascii=False
    )


def _stringify(value: Any) -> str:
    """Return ``value`` as a string; ``None`` becomes ``""``."""
    if value is None:
        return ""
    return str(value)


def _cell(row: list[Any], idx: int) -> str:
    """Return ``row[idx]`` as a string, tolerating short rows and ``None``."""
    if idx >= len(row) or row[idx] is None:
        return ""
    return str(row[idx])


def _coerce_turn(value: str) -> int:
    """Best-effort parse of a turn cell into an int; defaults to ``1``."""
    text = value.strip()
    if not text:
        return 1
    try:
        return int(float(text))
    except ValueError:
        return 1
