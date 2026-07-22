"""Parse a conversation ``.xlsx`` file into raw message rows.

The source spreadsheet stores a conversation as **one message per row** (a user
or model message with, optionally, a turn number). This module reads such a
file with ``openpyxl`` and projects it into a JSON-serializable list of message
dicts with the keys ``turn``, ``role`` and ``content`` (plus an optional
``retrieved_context_source`` — the **raw** ``retrieved_context`` cell, a ranked list
of source references the retrieval pipeline later fetches/extracts — and/or
``expected_output`` when the sheet carries those columns) — the raw conversation
shape that ``ScenarioResult.raw_conversation`` is documented to hold. Turn rows
for evaluation are derived from this projection in a later step; this module
only parses, and does **not** interpret the ``retrieved_context`` cell.

Header names are matched case-insensitively against a small alias table so the
parser tolerates either English or Spanish files (conversation content is in
Spanish). Callers with an unusual layout can pass an explicit ``columns``
mapping to override detection.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

# Header aliases (normalized: stripped + lowercased) -> canonical field name.
_HEADER_ALIASES: dict[str, str] = {
    "turn": "turn",
    "turno": "turn",
    "#": "turn",
    "role": "role",
    "rol": "role",
    "autor": "role",
    "content": "content",
    "contenido": "content",
    "mensaje": "content",
    "texto": "content",
    "retrieved context": "retrieved_context",
    "retrieved_context": "retrieved_context",
    "retrieval context": "retrieved_context",
    "retrieval_context": "retrieved_context",
    "contexto recuperado": "retrieved_context",
    "contexto": "retrieved_context",
    "expected output": "expected_output",
    "expected_output": "expected_output",
    "expected": "expected_output",
    "reference": "expected_output",
    "respuesta esperada": "expected_output",
    "salida esperada": "expected_output",
    "referencia": "expected_output",
    "esperado": "expected_output",
}

# Role value aliases (normalized) -> canonical role.
_ROLE_ALIASES: dict[str, str] = {
    "user": "user",
    "usuario": "user",
    "model": "model",
    "modelo": "model",
    "assistant": "model",
    "asistente": "model",
}


def _normalize(value: Any) -> str:
    """Strip and lowercase a cell/header value; ``None`` becomes ``""``."""
    if value is None:
        return ""
    return str(value).strip().lower()


def _map_columns(header: list[Any], override: dict[str, str] | None) -> dict[str, int]:
    """Map canonical field names to their 0-based column index in ``header``.

    ``override`` maps a canonical field (``turn``/``role``/``content``) to the
    header label as it literally appears in the sheet; it wins over aliases.
    Raises ``ValueError`` if ``role`` or ``content`` cannot be located.
    """
    normalized = [_normalize(cell) for cell in header]
    columns: dict[str, int] = {}

    if override:
        for field, label in override.items():
            target = _normalize(label)
            for idx, name in enumerate(normalized):
                if name == target:
                    columns[field] = idx
                    break

    for idx, name in enumerate(normalized):
        field = _HEADER_ALIASES.get(name)
        if field is not None and field not in columns:
            columns[field] = idx

    missing = [field for field in ("role", "content") if field not in columns]
    if missing:
        raise ValueError(
            f"No se encontraron columnas requeridas {missing} en el encabezado "
            f"{[c for c in header]!r}."
        )
    return columns


def parse_conversation(
    path: str | Path,
    *,
    columns: dict[str, str] | None = None,
) -> list[dict[str, Any]]:
    """Parse an ``.xlsx`` conversation file into raw message dicts.

    Returns a list of ``{"turn": int, "role": str, "content": str}`` in sheet
    order; each dict also carries ``retrieved_context_source`` (the raw context
    cell string, interpreted later by the retrieval pipeline) and/or
    ``expected_output`` (a string) when the sheet has those columns. Roles are
    normalized to canonical values
    (``user``/``model``) when recognized, otherwise passed through normalized.
    When the sheet has no turn column, turn numbers are derived: each ``user``
    message that follows a non-user message starts a new turn, so a user+model
    pair shares one number.

    ``columns`` optionally maps a canonical field (``turn``/``role``/``content``/
    ``retrieved_context``/``expected_output``) to the exact header label in the
    sheet, overriding alias detection.

    Raises ``FileNotFoundError`` if ``path`` does not exist and ``ValueError``
    if the required ``role``/``content`` columns cannot be found.
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"No existe el archivo de conversación: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            return []  # No worksheets in the workbook.
        rows = sheet.iter_rows(values_only=True)
        try:
            header = list(next(rows))
        except StopIteration:
            return []  # Empty sheet — no header, no messages.

        col = _map_columns(header, columns)
        has_turn_col = "turn" in col
        has_context_col = "retrieved_context" in col
        has_expected_col = "expected_output" in col

        messages: list[dict[str, Any]] = []
        derived_turn = 0
        prev_role: str | None = None

        for raw in rows:
            row = list(raw)
            content = _cell(row, col["content"]).strip()
            role_raw = _cell(row, col["role"]).strip()
            if not content and not role_raw:
                continue  # Fully-empty row.

            role = _ROLE_ALIASES.get(role_raw.lower(), role_raw.lower())

            if has_turn_col:
                turn = _coerce_turn(_cell(row, col["turn"]))
            else:
                if role == "user" and prev_role != "user":
                    derived_turn += 1
                elif derived_turn == 0:
                    derived_turn = 1
                turn = derived_turn

            message: dict[str, Any] = {"turn": turn, "role": role, "content": content}
            if has_context_col:
                # Store the raw cell verbatim; the retrieval pipeline interprets it later.
                message["retrieved_context_source"] = _cell(
                    row, col["retrieved_context"]
                ).strip()
            if has_expected_col:
                message["expected_output"] = _cell(
                    row, col["expected_output"]
                ).strip()
            messages.append(message)
            prev_role = role

        return messages
    finally:
        workbook.close()


def parse_conversation_json(
    path: str | Path,
    *,
    columns: dict[str, str] | None = None,
) -> str:
    """Parse ``path`` and return the raw conversation as a JSON string.

    Uses ``ensure_ascii=False`` so Spanish accents are preserved verbatim.
    """
    return json.dumps(
        parse_conversation(path, columns=columns), ensure_ascii=False
    )


def _cell(row: list[Any], idx: int) -> str:
    """Return ``row[idx]`` as a string, tolerating short rows and ``None``."""
    if idx >= len(row) or row[idx] is None:
        return ""
    return str(row[idx])


def _coerce_turn(value: str) -> int:
    """Best-effort parse of a turn cell into an int; defaults to ``1``."""
    text = value.strip()
    if not text:
        return 1
    try:
        return int(float(text))
    except ValueError:
        return 1
